"""Plan de labor e informe de sesión (Fase 2)."""

import csv
import zipfile
from datetime import date

from conftest import FIXTURES, ClienteFalso, leer

from informe_expedientes import plan_labor
from informe_expedientes.identificador import IdExpediente, encontrar
from informe_expedientes.sesion import OpcionesSesion, generar_sesion, resumen

HOY = date(2026, 10, 31)
B = "https://www.senado.gob.ar"


def _plan(nombre: str) -> plan_labor.PlanDeLabor:
    return plan_labor.leer(plan_labor.texto_de_pdf((FIXTURES / nombre).read_bytes()), nombre)


def _resumen_asuntos(plan):
    return [(a.orden, a.titulo, [r.clave for r in a.referencias]) for a in plan.asuntos]


def test_expedientes_en_texto_libre_y_listas_abreviadas():
    assert [i.corto for _p, i in encontrar("Tratamiento conjunto S-285, 637 y 638/24 y C.D. 14/24")] == [
        "S-285/24", "S-637/24", "S-638/24", "CD-14/24"]
    assert [i.corto for _p, i in encontrar("los exptes. S. 285-637-638-984- 1092 Y 1560/24")] == [
        "S-285/24", "S-637/24", "S-638/24", "S-984/24", "S-1092/24", "S-1560/24"]
    assert [i.corto for _p, i in encontrar("1370-S-2020 y PE 107/24; P.E. 186/25")] == [
        "S-1370/20", "PE-107/24", "PE-186/25"]
    assert encontrar("Ley 27.802, decreto 879/24, Art. 81 y OD 409/26") == []


def test_plan_del_28_de_octubre():
    plan = _plan("plan_labor_2026-10-28.pdf")
    assert (plan.fecha_reunion, plan.fecha_sesion, plan.hora) == (date(2026, 10, 28), date(2026, 10, 29), "14:00")
    assert _resumen_asuntos(plan) == [
        (1, "Jura de la Sra. Senadora Claudia López", []),
        (2, "RIIR", ["OD 413/26"]),
        (3, "Mercado de carbono", ["OD 366/26"]),
        (4, "Salud mental", ["OD 409/26"]),
        (5, "Pictogramas", ["OD 241/26"]),
        (6, "Transferencia de inmueble (Puerto Madryn)", ["OD 132/26"]),
        (7, "Educación emocional", ["OD 370/26"]),
        (8, "Proyecto de resolución: Papa León XIV", []),
        (9, "Proyecto de resolución: Citación al Jefe de Gabinete de Ministros", []),
    ]
    assert plan.asuntos[3].anexo and plan.asuntos[7].seccion == "Sobre tablas"


def test_listas_abreviadas_de_ordenes_del_dia():
    plan = _plan("plan_labor_2026-09-17.pdf")
    asuntos = _resumen_asuntos(plan)
    assert asuntos[1] == (2, "Acuerdos", ["OD 338/26", "OD 353/26", "OD 354/26"])        # «338; 353; 354»
    assert asuntos[2] == (3, "Traspaso Justicia Nacional a CABA", ["OD 360/26", "OD 361/26"])
    assert asuntos[4][2] == [f"OD {n}/26" for n in (340, 341, 342, *range(344, 353))]    # «340 a 342 y 344 a 352»
    assert asuntos[5][1] == "Biocombustibles"                                          # «OD 362 (con anexo)»
    # Sin título en la lista de temas: lo toma del grupo de la lista de impresos.
    assert asuntos[6][1:] == ("Reconocimientos héroes nacionales", ["OD 337/26", "OD 363/26", "OD 365/26"])
    assert asuntos[7][1] == "Convenio de creación de parque interjurisdiccional"


def test_rangos_con_anio_e_inclusive():
    plan = _plan("plan_labor_2026-07-08.pdf")
    acuerdos = plan.asuntos[1]
    assert acuerdos.ordenes_del_dia == ([f"{n}/26" for n in range(22, 28)] + [f"{n}/26" for n in range(123, 130)]
                                        + [f"{n}/26" for n in range(198, 220)])
    assert _resumen_asuntos(plan)[2:] == [(3, "Inviolabilidad de la Propiedad Privada", ["OD 104/26"]),
                                          (4, "Tucumán Capital simbólica", ["OD 152/26"])]
    assert plan_labor.referencias("● 0D 206/26", 2026, sueltas=True)[0][1].od == "206/26"


def test_texto_pegado_a_mano():
    plan = plan_labor.leer("1. OD 413/26 - RIIR\n2. Tratamiento de los expedientes S-285, 637 y 638/24\n"
                           "3. CD-14/24 (financiamiento universitario)\nHomenaje a los caídos", anio=2026)
    assert _resumen_asuntos(plan) == [(1, "RIIR", ["OD 413/26"]),
                                      (2, "Tratamiento de los expedientes", ["S-285/24", "S-637/24", "S-638/24"]),
                                      (3, "financiamiento universitario", ["CD-14/24"]),
                                      (4, "Homenaje a los caídos", [])]


def test_texto_legible_da_los_mismos_asuntos():
    texto = plan_labor.texto_de_pdf((FIXTURES / "plan_labor_2026-09-17.pdf").read_bytes())
    legible = plan_labor.texto_legible(texto)
    assert "\n2. OD de acuerdos: 338; 353; 354\n" in legible
    assert _resumen_asuntos(plan_labor.leer(legible)) == _resumen_asuntos(plan_labor.leer(texto))


def test_plan_en_word_y_en_texto(tmp_path):
    from docx import Document
    texto = plan_labor.texto_legible(plan_labor.texto_de_pdf((FIXTURES / "plan_labor_2026-10-28.pdf").read_bytes()))
    doc = Document()
    for linea in texto.splitlines():
        doc.add_paragraph(linea)
    doc.save(tmp_path / "plan.docx")
    desde_word = plan_labor.leer(plan_labor.texto_de_archivo(tmp_path / "plan.docx"))
    assert _resumen_asuntos(desde_word) == _resumen_asuntos(_plan("plan_labor_2026-10-28.pdf"))
    (tmp_path / "plan.txt").write_bytes("4. OD 409/26 y Anexo - Salud mental\n".encode("cp1252"))
    desde_txt = plan_labor.leer(plan_labor.texto_de_archivo(tmp_path / "plan.txt"), anio=2026)
    assert _resumen_asuntos(desde_txt) == [(1, "Salud mental", ["OD 409/26"])]


def test_plenario_de_labor_publicado():
    planes = plan_labor.parsear_plenario(leer("plenario_labor.html"))
    assert planes[0].fecha_reunion == date(2026, 10, 28) and planes[0].fecha_sesion == date(2026, 10, 29)
    assert planes[0].url.endswith("Acta%20de%20labor%20parlamentaria%20%2028-10-2026.pdf/download")
    sin_sesion = next(x for x in planes if x.fecha_reunion == date(2026, 2, 24))
    assert sin_sesion.fecha_sesion is None and "Plan%20de%20Labor%2024_02_2026" in sin_sesion.url
    assert planes[0].etiqueta == "sesión del 29/10/2026 (reunión del 28/10/2026)"


def test_expedientes_de_las_ordenes_del_dia():
    assert plan_labor.expedientes_de_od(leer("od_2026_409.html")) == [IdExpediente(99, 2026, "PE", "PL")]
    assert [i.corto for i in plan_labor.expedientes_de_od(leer("od_2026_366.html"))] == [
        "S-718/26", "S-808/26", "S-865/26"]
    assert plan_labor.url_od("409/26") == f"{B}/parlamentario/parlamentaria/ordenDelDiaResultadoLink/2026/409"


def test_informe_de_sesion_completo(tmp_path):
    res = generar_sesion(_plan("plan_labor_2026-10-28.pdf"), OpcionesSesion(carpeta_salida=tmp_path), ClienteFalso(),
                         hoy=HOY)
    assert resumen(res) == ("El plan de labor para la sesión del 29/10/2026 acordó NUEVE (9) asuntos: SEIS (6) con "
                            "orden del día o expediente y TRES (3) sin número (juras, acuerdos o asuntos sobre "
                            "tablas). Se consultaron las fichas de NUEVE (9) expedientes.")
    assert [x.corto for x in res.temas[2].expedientes] == ["S-718/26", "S-808/26", "S-865/26"]
    assert res.temas[3].resultados[0].ficha.id.corto == "PE-99/26"
    assert set(res.archivos) == {"docx", "html", "xlsx", "csv"}

    with open(res.archivos["csv"], encoding="utf-8-sig", newline="") as f:
        filas = list(csv.reader(f, delimiter=";"))
    assert filas[0] == ["Fecha de sesión", "Orden", "Expediente", "OD", "Título", "Tipo", "Comisiones"]
    assert filas[1] == ["29/10/2026", "1", "", "", "Jura de la Sra. Senadora Claudia López", "", ""]
    assert filas[3] == ["29/10/2026", "3", "S-718/26, S-808/26, S-865/26", "366/26", "Mercado de carbono",
                        "Proyecto de ley", "Legislación General / Ambiente y Desarrollo Sustentable / "
                        "Presupuesto y Hacienda"]
    assert filas[6][2:6] == ["S-695/26, S-700/26", "132/26", "Transferencia de inmueble (Puerto Madryn)",
                             "Proyecto de ley, Comunicaciones de senadores"]
    assert len(filas) == 10
    assert res.archivos["csv"].read_bytes().startswith(b"\xef\xbb\xbf")

    html = res.archivos["html"].read_text(encoding="utf-8")
    assert '<a href="#tema-4">Salud mental</a>' in html and 'id="tema-4"' in html
    assert html.count('class="ficha"') == 9 and "Asuntos en el orden de tratamiento" in html
    assert "#2E6B5F" in html
    with zipfile.ZipFile(res.archivos["docx"]) as z:
        documento = z.read("word/document.xml").decode("utf-8")
    assert "Asuntos en el orden de tratamiento" in documento and "Educación emocional" in documento
    assert documento.count("Situación actual") == 9
    from openpyxl import load_workbook
    hoja = load_workbook(res.archivos["xlsx"])["Asuntos"]
    assert hoja["A4"].value == "Orden" and hoja.max_row == 4 + 12           # 9 fichas y 3 asuntos sin número


def test_od_que_no_se_puede_abrir(tmp_path):
    plan = plan_labor.leer("1. OD 999/26 - Inexistente", anio=2026)
    res = generar_sesion(plan, OpcionesSesion(carpeta_salida=tmp_path, formatos=("csv",)), ClienteFalso(), hoy=HOY)
    assert res.temas[0].resultados == [] and "999/26" in res.advertencias[0]
    assert res.archivos["csv"].exists()


def test_linea_de_comandos_de_sesion(tmp_path, monkeypatch, capsys):
    import app
    monkeypatch.setattr(app, "Cliente", lambda *a, **k: ClienteFalso())
    archivo = str(FIXTURES / "plan_labor_2026-09-17.pdf")
    assert app.linea_de_comandos(["--sesion", archivo, "--solo-asuntos"]) == 0
    salida = capsys.readouterr().out
    assert " 5. Cónsules y Vicecónsules Honorarios  [OD 340/26, OD 341/26" in salida
    codigo = app.linea_de_comandos(["--sesion", "ultima", "--salida", str(tmp_path), "--nombre", "sesion",
                                    "--formatos", "html,csv"])
    salida = capsys.readouterr().out
    assert codigo == 0 and "Plan de labor publicado: sesión del 29/10/2026" in salida
    assert "      PE-99/26: Media sanción del Senado" in salida
    assert (tmp_path / "sesion.html").exists() and (tmp_path / "sesion.csv").exists()
    assert app.linea_de_comandos(["--sesion", "01/01/2001", "--solo-asuntos"]) == 2
    assert app.linea_de_comandos(["--sesion", "no-existe.pdf"]) == 2
