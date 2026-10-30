"""Informe en Word (.docx), con enlaces que se pueden abrir con Ctrl+clic."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import __version__
from .contenido import (
    Linea, alcance_busqueda, bloques_resolucion, comisiones_texto, datos_diputados, datos_generales,
    descripcion, descripcion_votacion, enlaces_votacion, nota_votacion, resultado_votacion, texto_votaciones,
    tipo_votacion, totales_votacion, votaciones_con_bloques,
)
from .modelos import Resultado
from .util import fmt_fecha, nombre_comision

C = {
    "tinta": "1F2B28", "suave": "55645F", "cupula": "2E6B5F", "cupula_oscura": "1F4E45",
    "cupula_tenue": "E4EFEB", "piedra": "EEF2F0", "linea": "CBD5D1", "blanco": "FFFFFF",
    "lacre": "8C2F2F", "lacre_tenue": "F6E7E5", "ambar": "8A5A10", "ambar_tenue": "F7EEDC",
}
SERIF = "Georgia"
SANS = "Segoe UI"
ANCHO_UTIL = 16.6  # cm (A4 con márgenes de 2,2 cm)

_FONDO_NIVEL = {
    "con_dictamen": ("ambar_tenue", "ambar"),
    "vetado": ("lacre_tenue", "lacre"), "caducado": ("lacre_tenue", "lacre"),
    "rechazado": ("lacre_tenue", "lacre"), "archivado": ("lacre_tenue", "lacre"),
}


# ======================================================================= XML

def _fuente_estilo(estilo, nombre: str, tamano: float | None = None, color: str | None = None,
                   negrita: bool | None = None) -> None:
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(atributo), None)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atributo), nombre)
    if tamano:
        estilo.font.size = Pt(tamano)
    if color:
        estilo.font.color.rgb = RGBColor.from_string(color)
        nodo = rpr.find(qn("w:color"))
        if nodo is not None:
            for atributo in ("w:themeColor", "w:themeShade", "w:themeTint"):
                nodo.attrib.pop(qn(atributo), None)
    if negrita is not None:
        estilo.font.bold = negrita


def _sombrear(celda, color: str) -> None:
    tcpr = celda._tc.get_or_add_tcPr()
    for viejo in tcpr.findall(qn("w:shd")):
        tcpr.remove(viejo)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), C.get(color, color))
    tcpr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _borde_celda(celda, lado: str, color: str, grosor: int = 4) -> None:
    tcpr = celda._tc.get_or_add_tcPr()
    bordes = tcpr.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        tcpr.insert_element_before(bordes, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                                   "w:tcFitText", "w:vAlign", "w:hideMark")
    elemento = OxmlElement(f"w:{lado}")
    elemento.set(qn("w:val"), "single")
    elemento.set(qn("w:sz"), str(grosor))
    elemento.set(qn("w:space"), "0")
    elemento.set(qn("w:color"), C.get(color, color))
    bordes.append(elemento)


def _bordes_tabla(tabla, color: str = "linea", verticales: bool = False, grosor: int = 4) -> None:
    tblpr = tabla._tbl.tblPr
    for viejo in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(viejo)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        if lado in ("left", "right", "insideV") and not verticales:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(grosor))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), C.get(color, color))
        bordes.append(e)
    tblpr.insert_element_before(bordes, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription")


def _margenes_celdas(tabla, arriba=60, abajo=60, izquierda=100, derecha=100) -> None:
    tblpr = tabla._tbl.tblPr
    margenes = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", arriba), ("left", izquierda), ("bottom", abajo), ("right", derecha)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(valor))
        e.set(qn("w:type"), "dxa")
        margenes.append(e)
    tblpr.insert_element_before(margenes, "w:tblLook", "w:tblCaption", "w:tblDescription")


def _fila_encabezado(fila) -> None:
    trpr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def _no_partir(fila) -> None:
    trpr = fila._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _hipervinculo(parrafo, texto: str, url: str, tamano: float | None = None, negrita: bool = False) -> None:
    r_id = parrafo.part.relate_to(url, RT.HYPERLINK, is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), r_id)
    enlace.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if negrita:
        rpr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), C["cupula_oscura"])
    rpr.append(color)
    if tamano:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(tamano * 2)))
        rpr.append(sz)
    subrayado = OxmlElement("w:u")
    subrayado.set(qn("w:val"), "single")
    rpr.append(subrayado)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    enlace.append(run)
    parrafo._p.append(enlace)


def _campo(parrafo, instruccion: str, tamano: float = 8) -> None:
    def run_con(elemento):
        r = parrafo.add_run()
        r.font.size = Pt(tamano)
        r.font.color.rgb = RGBColor.from_string(C["suave"])
        r._r.append(elemento)
        return r

    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    run_con(inicio)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "
    run_con(instr)
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    run_con(separador)
    r = parrafo.add_run("1")
    r.font.size = Pt(tamano)
    r.font.color.rgb = RGBColor.from_string(C["suave"])
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run_con(fin)


# ================================================================= escritura

class _Escritor:
    def __init__(self, resultado, encabezado: str | None = None, apaisado: bool = False):
        self.r = resultado
        self.apaisado = apaisado
        self.ancho = 25.3 if apaisado else ANCHO_UTIL
        self.encabezado = encabezado or f"H. Senado de la Nación — Informe de expediente {resultado.ficha.id.corto}"
        self.doc = Document()
        self._preparar()

    # ---- estilos y página
    def _preparar(self) -> None:
        seccion = self.doc.sections[0]
        if self.apaisado:
            seccion.orientation = WD_ORIENT.LANDSCAPE
            seccion.page_width, seccion.page_height = Cm(29.7), Cm(21)
        else:
            seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
        for lado in ("left_margin", "right_margin"):
            setattr(seccion, lado, Cm(2.2))
        seccion.top_margin, seccion.bottom_margin = Cm(2.0), Cm(2.0)
        seccion.header_distance, seccion.footer_distance = Cm(1.0), Cm(1.0)

        estilos = self.doc.styles
        normal = estilos["Normal"]
        _fuente_estilo(normal, SANS, 10, C["tinta"])
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = 1.15
        for nombre, tamano in (("Heading 1", 14), ("Heading 2", 11.5)):
            estilo = estilos[nombre]
            _fuente_estilo(estilo, SERIF, tamano, C["cupula_oscura"], negrita=False)
            estilo.paragraph_format.space_before = Pt(16 if nombre == "Heading 1" else 10)
            estilo.paragraph_format.space_after = Pt(6)
            estilo.paragraph_format.keep_with_next = True
        try:
            _fuente_estilo(estilos["List Bullet"], SANS, 10, C["tinta"])
        except KeyError:
            pass

        seccion.different_first_page_header_footer = True
        encabezado = seccion.header.paragraphs[0]
        run = encabezado.add_run(self.encabezado)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(C["suave"])
        for pie in (seccion.footer.paragraphs[0], seccion.first_page_footer.paragraphs[0]):
            run = pie.add_run(f"Generado el {self.r.generado:%d/%m/%Y %H:%M} h con datos públicos de senado.gob.ar")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(C["suave"])
            pie.paragraph_format.tab_stops.add_tab_stop(Cm(self.ancho), WD_TAB_ALIGNMENT.RIGHT)
            r = pie.add_run("\tPágina ")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(C["suave"])
            _campo(pie, "PAGE")
            r = pie.add_run(" de ")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(C["suave"])
            _campo(pie, "NUMPAGES")

    # ---- piezas
    def parrafo(self, texto: str = "", tamano: float | None = None, color: str | None = None,
                fuente: str | None = None, negrita: bool = False, cursiva: bool = False,
                antes: float | None = None, despues: float | None = None):
        p = self.doc.add_paragraph()
        if texto:
            self._run(p, texto, tamano, color, fuente, negrita, cursiva)
        if antes is not None:
            p.paragraph_format.space_before = Pt(antes)
        if despues is not None:
            p.paragraph_format.space_after = Pt(despues)
        return p

    @staticmethod
    def _run(p, texto, tamano=None, color=None, fuente=None, negrita=False, cursiva=False):
        run = p.add_run(texto)
        if tamano:
            run.font.size = Pt(tamano)
        if color:
            run.font.color.rgb = RGBColor.from_string(C.get(color, color))
        if fuente:
            run.font.name = fuente
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), fuente)
        run.bold = negrita
        run.italic = cursiva
        return run

    def linea(self, p, linea: Linea, tamano: float | None = None) -> None:
        for fragmento in linea:
            if fragmento.url:
                _hipervinculo(p, fragmento.texto, fragmento.url, tamano)
            else:
                self._run(p, fragmento.texto, tamano)

    def titulo(self, texto: str) -> None:
        self.doc.add_heading(texto, level=1)

    def tabla(self, encabezados: list[str], anchos: list[float]):
        t = self.doc.add_table(rows=1, cols=len(encabezados))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        _bordes_tabla(t)
        _margenes_celdas(t)
        fila = t.rows[0]
        _fila_encabezado(fila)
        for i, (texto, ancho) in enumerate(zip(encabezados, anchos)):
            celda = fila.cells[i]
            celda.width = Cm(ancho)
            _sombrear(celda, "piedra")
            p = celda.paragraphs[0]
            self._run(p, texto, 9, "suave", negrita=True)
        for i, ancho in enumerate(anchos):
            t.columns[i].width = Cm(ancho)
        return t

    def fila(self, tabla, anchos: list[float]):
        fila = tabla.add_row()
        _no_partir(fila)
        for i, ancho in enumerate(anchos):
            fila.cells[i].width = Cm(ancho)
        return fila

    def pares(self, filas: list[tuple[str, Linea]]) -> None:
        anchos = [5.2, self.ancho - 5.2]
        t = self.doc.add_table(rows=0, cols=2)
        t.autofit = False
        _bordes_tabla(t)
        _margenes_celdas(t, 50, 50, 60, 100)
        for etiqueta, linea in filas:
            fila = self.fila(t, anchos)
            self._run(fila.cells[0].paragraphs[0], etiqueta, 9.5, "suave")
            self.linea(fila.cells[1].paragraphs[0], linea)
        for i, ancho in enumerate(anchos):
            t.columns[i].width = Cm(ancho)

    # ---- documento
    def escribir(self, destino: Path, ejemplo: bool = False) -> Path:
        r, f = self.r, self.r.ficha

        p = self.parrafo("H. Senado de la Nación", 9, "suave", despues=0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(self.ancho), WD_TAB_ALIGNMENT.RIGHT)
        self._run(p, f"\tInforme del {r.generado:%d/%m/%Y, %H:%M} h", 9, "suave")
        if ejemplo:
            p = self.parrafo("Ejemplo generado con páginas de prueba. ", 9, "lacre", negrita=True, antes=4)
            self._run(p, "No reemplaza una consulta real.", 9, "lacre")
        self.parrafo(f.id.corto, 32, "tinta", SERIF, antes=14, despues=2)
        self.parrafo(descripcion(f), 11, "suave", despues=8)
        if f.extracto:
            self.parrafo(f.extracto, 11, "tinta", SERIF, despues=8)
        p = self.parrafo(despues=6)
        _hipervinculo(p, "Ficha en el sitio del Senado", f.url, 9.5)
        if f.texto_original:
            self._run(p, "      ")
            _hipervinculo(p, "Texto original (PDF)", f.texto_original.url, 9.5)
        for d in f.textos_definitivos:
            self._run(p, "      ")
            _hipervinculo(p, d.etiqueta, d.url, 9.5)

        self._situacion()

        self.titulo("Datos del expediente")
        self.pares(datos_generales(f))

        dip = datos_diputados(f)
        if dip:
            self.titulo("Etapa en la Cámara de Diputados")
            self.pares(dip)

        if f.giros:
            self.titulo("Giros a comisiones")
            anchos = [8.6, 2.6, 2.7, 2.7]
            t = self.tabla(["Comisión", "Orden de giro", "Ingreso", "Egreso"], anchos)
            for g in f.giros:
                fila = self.fila(t, anchos)
                valores = [nombre_comision(g.comision), str(g.orden or "—"), fmt_fecha(g.ingreso),
                           fmt_fecha(g.egreso, "en curso" if g.ingreso else "—")]
                for i, v in enumerate(valores):
                    self._run(fila.cells[i].paragraphs[0], v, 9.5)

        self._reuniones()

        if r.convocadas:
            self.titulo("Próximas reuniones convocadas")
            anchos = [2.6, 1.6, 8.2, 4.2]
            t = self.tabla(["Fecha", "Hora", "Reunión", "Lugar"], anchos)
            for c in r.convocadas:
                fila = self.fila(t, anchos)
                self._run(fila.cells[0].paragraphs[0], fmt_fecha(c.fecha, c.fecha_texto), 9.5)
                self._run(fila.cells[1].paragraphs[0], c.hora or "—", 9.5)
                _hipervinculo(fila.cells[2].paragraphs[0], c.titulo, c.url, 9.5)
                self._run(fila.cells[2].add_paragraph(), c.comision, 8.5, "suave")
                self._run(fila.cells[3].paragraphs[0], c.lugar or "—", 9.5)

        if f.ordenes_del_dia:
            self.titulo("Órdenes del Día del Senado")
            anchos = [3.0, 2.6, 6.8, 4.2]
            t = self.tabla(["Número", "Fecha", "Estado", "Texto"], anchos)
            for od in f.ordenes_del_dia:
                fila = self.fila(t, anchos)
                if od.url:
                    _hipervinculo(fila.cells[0].paragraphs[0], f"N° {od.numero}", od.url, 9.5)
                else:
                    self._run(fila.cells[0].paragraphs[0], f"N° {od.numero}", 9.5)
                self._run(fila.cells[1].paragraphs[0], fmt_fecha(od.fecha), 9.5)
                self._run(fila.cells[2].paragraphs[0], od.estado.capitalize(), 9.5)
                if od.url_pdf:
                    _hipervinculo(fila.cells[3].paragraphs[0], "Dictamen (PDF)", od.url_pdf, 9.5)
                else:
                    self._run(fila.cells[3].paragraphs[0], "—", 9.5)

        if r.votaciones_consultadas:
            self._votaciones()

        bloques = bloques_resolucion(f)
        if bloques:
            self.titulo("Sanción y resoluciones")
            for titulo, campos in bloques:
                self.doc.add_heading(titulo, level=2)
                self.pares([(k, [_frag(v)]) for k, v in campos])

        observaciones = f.observaciones + f.anotaciones + [
            h for h in f.hitos if not any(k in h.upper() for k in ("ARCHIVO", "CADUC", "PREFERENCIA", "PARA:"))]
        if observaciones:
            self.titulo("Observaciones de la ficha")
            for o in observaciones:
                self.doc.add_paragraph(o, style="List Bullet")

        if r.notas_usuario:
            self.titulo("Notas")
            for linea in r.notas_usuario.splitlines():
                self.parrafo(linea, 10.5, fuente=SERIF)

        p = self.parrafo(antes=22)
        self._run(p, "Informe generado automáticamente con datos públicos del sitio del H. Senado de la Nación. "
                     "Antes de usarlo en un trámite oficial, verificá los datos en la ", 8.5, "suave")
        _hipervinculo(p, "ficha del expediente", f.url, 8.5)
        self._run(p, f". Programa Informe de Expedientes, versión {__version__}.", 8.5, "suave")
        if r.advertencias:
            self.parrafo("Durante la consulta hubo avisos:", 8.5, "suave", antes=4)
            for a in r.advertencias:
                p = self.doc.add_paragraph(style="List Bullet")
                self._run(p, a, 8.5, "suave")

        self.doc.core_properties.title = f"Informe del expediente {f.id.corto}"
        self.doc.core_properties.subject = f.extracto[:250]
        self.doc.core_properties.author = "Informe de Expedientes"
        self.doc.save(str(destino))
        return Path(destino)

    def _situacion(self) -> None:
        s = self.r.situacion
        fondo, borde = _FONDO_NIVEL.get(s.nivel, ("cupula_tenue", "cupula"))
        caja = self.doc.add_table(rows=1, cols=1)
        caja.autofit = False
        _bordes_tabla(caja, color=fondo)
        _margenes_celdas(caja, 160, 140, 260, 200)
        celda = caja.rows[0].cells[0]
        celda.width = Cm(self.ancho)
        _sombrear(celda, fondo)
        _borde_celda(celda, "left", borde, 36)
        p = celda.paragraphs[0]
        self._run(p, "Situación actual", 9, "suave")
        p = celda.add_paragraph()
        self._run(p, s.titulo, 15, "tinta", SERIF)
        p.paragraph_format.space_after = Pt(4)
        for d in s.detalle:
            p = celda.add_paragraph()
            self._run(p, "•  " + d, 9.5)
            p.paragraph_format.left_indent = Cm(0.3)
            p.paragraph_format.space_after = Pt(1)
        if s.caducidad:
            c = s.caducidad
            p = celda.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            etiqueta = "Caducidad registrada" if c.registrada else "Caducidad estimada"
            color = "lacre" if (c.vencida and not c.registrada) else "tinta"
            self._run(p, f"{etiqueta}: {fmt_fecha(c.fecha)}. ", 9, color, negrita=True)
            aviso = " El plazo estimado ya venció: verificar." if c.vencida and not c.registrada else ""
            self._run(p, c.fundamento + aviso, 8.5, "suave")

        if s.pasos:
            espacio = self.parrafo(despues=0)
            espacio.paragraph_format.space_before = Pt(2)
            ancho = self.ancho / len(s.pasos)
            t = self.doc.add_table(rows=1, cols=len(s.pasos))
            t.autofit = False
            _bordes_tabla(t, color="blanco", verticales=True, grosor=18)
            _margenes_celdas(t, 70, 70, 90, 70)
            estilos = {
                "hecho": ("cupula_tenue", "tinta", "✓ "),
                "actual": ("cupula", "blanco", "● "),
                "detenido": ("lacre_tenue", "lacre", "× "),
                "pendiente": ("piedra", "suave", ""),
            }
            for i, paso in enumerate(s.pasos):
                celda = t.rows[0].cells[i]
                celda.width = Cm(ancho)
                fondo, color, marca = estilos[paso.estado]
                _sombrear(celda, fondo)
                p = celda.paragraphs[0]
                self._run(p, f"{i + 1}. {marca}{paso.nombre}", 8.5, color, negrita=True)
                extra = " ".join(x for x in [fmt_fecha(paso.fecha, ""), paso.detalle] if x)
                if extra:
                    p2 = celda.add_paragraph()
                    self._run(p2, extra, 7.5, color)
            for i in range(len(s.pasos)):
                t.columns[i].width = Cm(ancho)

    def _votaciones(self) -> None:
        r = self.r
        self.titulo("Votaciones en el recinto")
        self.parrafo(texto_votaciones(r), 10)
        if r.votaciones:
            anchos = [2.0, 5.4, 2.2, 1.9, 1.0, 1.0, 1.0, 1.0, 1.1]
            t = self.tabla(["Fecha", "Votación", "Tipo", "Resultado", "Afirm.", "Neg.", "Abst.", "Aus.", "Acta"],
                           anchos)
            for v in r.votaciones:
                fila = self.fila(t, anchos)
                self._run(fila.cells[0].paragraphs[0], fmt_fecha(v.fecha), 9)
                self._run(fila.cells[1].paragraphs[0], descripcion_votacion(v), 9)
                self._run(fila.cells[2].paragraphs[0], tipo_votacion(v), 9)
                self._run(fila.cells[3].paragraphs[0], resultado_votacion(v), 9, "lacre" if v.rechazada else None)
                for i, numero in enumerate(totales_votacion(v), start=4):
                    p = fila.cells[i].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    self._run(p, numero, 9)
                enlaces = enlaces_votacion(v)
                if not enlaces:
                    self._run(fila.cells[8].paragraphs[0], "—", 9)
                for j, fr in enumerate(enlaces):
                    p = fila.cells[8].paragraphs[0] if j == 0 else fila.cells[8].add_paragraph()
                    _hipervinculo(p, fr.texto, fr.url, 8.5)
        for v in votaciones_con_bloques(r):
            self.doc.add_heading(f"Cómo votó cada bloque: {tipo_votacion(v).lower()}, {fmt_fecha(v.fecha)}", level=2)
            anchos = [7.0, 2.4, 2.4, 2.4, 2.4]
            t = self.tabla(["Bloque", "Afirmativos", "Negativos", "Abstenciones", "Ausentes"], anchos)
            for bloque, *numeros in v.por_bloque():
                fila = self.fila(t, anchos)
                self._run(fila.cells[0].paragraphs[0], bloque, 9)
                for i, n in enumerate(numeros, start=1):
                    p = fila.cells[i].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    self._run(p, str(n), 9)
            nota = nota_votacion(v)
            if nota:
                self.parrafo(nota, 8.5, "suave", antes=4)

    def _reuniones(self) -> None:
        r = self.r
        self.titulo("Reuniones de comisión en que se trató")
        if r.reuniones:
            anchos = [2.4, 7.4, 2.2, 4.6]
            t = self.tabla(["Fecha", "Reunión", "Tipo", "Documentos"], anchos)
            for reunion in r.reuniones:
                fila = self.fila(t, anchos)
                _hipervinculo(fila.cells[0].paragraphs[0], fmt_fecha(reunion.fecha), reunion.url, 9.5)
                _hipervinculo(fila.cells[1].paragraphs[0], reunion.titulo, reunion.url, 9.5)
                comisiones = comisiones_texto(reunion.comisiones)
                if comisiones:
                    self._run(fila.cells[1].add_paragraph(), comisiones, 8.5, "suave")
                self._run(fila.cells[2].paragraphs[0], reunion.tipo, 9, "suave")
                celda = fila.cells[3]
                documentos = [(d.etiqueta if len(d.etiqueta) > 3 else "Versión taquigráfica", d.url)
                              for d in reunion.versiones]
                documentos += [("Video", v) for v in reunion.videos]
                documentos += [(d.etiqueta, d.url) for d in reunion.otros_archivos]
                if not documentos:
                    self._run(celda.paragraphs[0], "sin archivos", 9, "suave", cursiva=True)
                for j, (texto, url) in enumerate(documentos):
                    p = celda.paragraphs[0] if j == 0 else celda.add_paragraph()
                    _hipervinculo(p, texto, url, 9)
        elif r.busqueda.realizada:
            self.parrafo("No se encontraron reuniones publicadas en que figure este expediente.",
                         10, "suave", cursiva=True)
        else:
            self.parrafo("No se buscaron reuniones de comisión para este informe.", 10, "suave", cursiva=True)
        self.parrafo(alcance_busqueda(r), 8.5, "suave", antes=4)


def _frag(texto: str):
    from .contenido import Fragmento
    return Fragmento(texto)


def escribir(resultado: Resultado, destino: Path, ejemplo: bool = False) -> Path:
    return _Escritor(resultado).escribir(Path(destino), ejemplo)
