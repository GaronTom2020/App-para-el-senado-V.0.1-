"""Plan de labor parlamentaria: lectura y asuntos en el orden de tratamiento.

El plan llega como archivo (Word, PDF o texto pegado) o se baja del sitio del Senado
(«Plenario de Labor Parlamentaria», un PDF por reunión). Su estructura, relevada en
docs/relevamiento-fase0.md, es: fecha de la reunión y de la sesión, «Se acordó el
tratamiento de los siguientes temas» (lista numerada en el orden de tratamiento),
«Órdenes del día impresos a considerar», preferencias y «Asuntos sobre tablas».

Casi siempre nombra órdenes del día, a veces en listas abreviadas («OD 340 a 342 y
344 a 352», «338; 353; 354», «OD 22/26 a 27/26 inclusive», incluso «0D 206/26»), y
a veces expedientes («S-285, 637 y 638/24»). Si el texto no tiene esa estructura (una
lista pegada a mano), se toma en el orden en que aparecen.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from urllib.parse import urljoin

from . import config
from .cliente import Cliente
from .identificador import _EN_TEXTO_NUMERO, _EN_TEXTO_ORIGEN, RE_URL_FICHA, IdExpediente, encontrar
from .paginas import sopa, texto as texto_de
from .util import anio_cuatro_digitos, limpiar, normalizar, parsear_fecha

URL_PLENARIO = config.BASE + "/parlamentario/parlamentaria/actas"
URL_OD = config.BASE + "/parlamentario/parlamentaria/ordenDelDiaResultadoLink/{anio}/{numero}"
TTL_PLENARIO = 3600
MAX_RANGO = 120          # un rango más largo que esto es casi seguro un error de lectura


class ErrorPlan(Exception):
    """No se pudo leer el plan de labor."""


@dataclass
class Referencia:
    clase: str                       # od · expediente
    od: str | None = None            # «409/26»
    ident: IdExpediente | None = None

    @property
    def clave(self) -> str:
        return f"OD {self.od}" if self.clase == "od" else self.ident.corto


@dataclass
class Asunto:
    orden: int
    titulo: str
    seccion: str                     # Temas acordados · Órdenes del día · Preferencias · Sobre tablas
    texto: str
    referencias: list[Referencia] = field(default_factory=list)
    anexo: bool = False

    @property
    def ordenes_del_dia(self) -> list[str]:
        return [r.od for r in self.referencias if r.clase == "od"]

    @property
    def expedientes(self) -> list[IdExpediente]:
        return [r.ident for r in self.referencias if r.clase == "expediente"]


@dataclass
class PlanDeLabor:
    asuntos: list[Asunto]
    fecha_reunion: date | None = None
    fecha_sesion: date | None = None
    hora: str = ""
    fuente: str = ""                 # nombre del archivo o dirección del PDF
    texto: str = ""


@dataclass
class PlanPublicado:
    fecha_reunion: date | None
    fecha_sesion: date | None
    url: str

    @property
    def etiqueta(self) -> str:
        sesion = f"sesión del {self.fecha_sesion:%d/%m/%Y}" if self.fecha_sesion else "sin fecha de sesión"
        reunion = f"reunión del {self.fecha_reunion:%d/%m/%Y}" if self.fecha_reunion else ""
        return f"{sesion} ({reunion})" if reunion else sesion


# ==================================================================== lectura

def texto_de_archivo(ruta: Path) -> str:
    """Texto de un plan en Word (.docx), PDF o texto plano."""
    ruta = Path(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo == ".docx":
        from docx import Document
        try:
            doc = Document(str(ruta))
        except Exception as e:  # noqa: BLE001 - archivo dañado o no es Word
            raise ErrorPlan(f"No se pudo abrir el Word ({type(e).__name__}).") from e
        lineas = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                lineas.append(" ".join(c.text for c in fila.cells))
        return "\n".join(lineas)
    if sufijo == ".pdf":
        return texto_de_pdf(ruta.read_bytes())
    if sufijo == ".doc":
        raise ErrorPlan("Los archivos .doc (Word 97-2003) no se pueden leer: guardalo como .docx o PDF.")
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return ruta.read_text(encoding=codificacion)
        except UnicodeDecodeError:
            continue
    raise ErrorPlan("No se pudo leer el archivo de texto.")


def texto_de_pdf(contenido: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ErrorPlan("Falta la biblioteca pypdf para leer el PDF.") from e
    try:
        lector = PdfReader(io.BytesIO(contenido))
        return "\n".join(p.extract_text() or "" for p in lector.pages)
    except Exception as e:  # noqa: BLE001 - PDF dañado o escaneado
        raise ErrorPlan(f"No se pudo leer el PDF ({type(e).__name__}).") from e


def texto_legible(texto: str) -> str:
    """El texto de un PDF (una palabra por renglón) en renglones: uno por tema, viñeta o sección."""
    t = limpiar(unicodedata.normalize("NFKC", texto or ""))
    n = _preparar(t)
    cortes = {0}
    for _nombre, patron in _SECCIONES:
        cortes.update(m.start() for m in re.finditer(patron, n))
    cortes.update(m.start() for m in re.finditer(r"[●•]|(?<![\w/.])\d{1,2}\s*[.)](?=\s)", t))
    cortes.update(m.start() for m in re.finditer(r"FECHA DE SESION|HORA DE SESION|CARACTER DE LA SESION", n))
    posiciones = sorted(cortes) + [len(t)]
    return "\n".join(limpiar(t[a:b]) for a, b in zip(posiciones, posiciones[1:]) if limpiar(t[a:b]))


def parsear_plenario(html: str) -> list[PlanPublicado]:
    """Tabla del Plenario de Labor Parlamentaria: reunión, sesión prevista y PDF, de la más nueva a la más vieja."""
    planes = []
    for tr in sopa(html).find_all("tr"):
        a = tr.find("a", href=re.compile(r"\.pdf/download$", re.I))
        tds = tr.find_all("td")
        if not a or len(tds) < 3:
            continue
        planes.append(PlanPublicado(parsear_fecha(texto_de(tds[0])), parsear_fecha(texto_de(tds[1])),
                                    urljoin(config.BASE, a["href"])))
    return planes


def planes_publicados(cliente: Cliente) -> list[PlanPublicado]:
    return parsear_plenario(cliente.get(URL_PLENARIO, ttl=TTL_PLENARIO).texto)


def plan_publicado(cliente: Cliente, cual: PlanPublicado) -> PlanDeLabor:
    """Baja y lee un plan publicado en el sitio. Hay que usar el enlace de la tabla: los nombres no siguen un patrón."""
    r = cliente.get(cual.url, binario=True)
    plan = leer(texto_de_pdf(r.contenido or b""), cual.url)
    plan.fecha_reunion = plan.fecha_reunion or cual.fecha_reunion
    plan.fecha_sesion = plan.fecha_sesion or cual.fecha_sesion
    return plan


# ================================================================ extracción

_SECCIONES = [
    ("Temas acordados", r"SE ACORDO (?:EL TRATAMIENTO DE |CONSIDERAR )?(?:LOS )?SIGUIENTES TEMAS\s*:?"),
    ("Órdenes del día", r"ORDENES DEL DIA IMPRESOS A CONSIDERAR"),
    ("Preferencias", r"PREFERENCIAS A TRATAR VOTADAS CON ANTERIORIDAD"),
    ("Sobre tablas", r"ASUNTOS SOBRE TABLAS ACORDADOS\s*:?"),
    ("Sobre tablas (a solicitar)", r"\bA SOLICITAR\s*:"),
    ("Preferencias a solicitar", r"PREFERENCIAS A SOLICITAR"),
    ("Fin", r"OBSERVACIONES Y ASUNTOS VARIOS|AUTORIDADES DE LA CAMARA"),
]

# «OD 340 a 342 y 344 a 352», «OD de acuerdos: 338; 353; 354», «OD 22/26 a 27/26 inclusive», «0D 206/26»
_NUM = r"\d{1,4}(?:\s*/\s*\d{2,4})?"
_RE_OD = re.compile(
    r"(?<![A-Z])(?:O\.?\s?D\.?|0D|ORDEN(?:ES)? DEL DIA)\s*(?:N[°º]\s*|NROS?\.?\s*)?(?:DE ACUERDOS?\s*:?\s*)?"
    rf"(?P<lista>{_NUM}(?:\s*(?:,|;|\bY\b|\bE\b|\bAL?\b|-)\s*{_NUM})*)(?:\s*INCLUSIVE)?")
_RE_OD_SUELTA = re.compile(r"[●•]\s*(?P<lista>\d{1,4}\s*/\s*\d{2,4})")       # «● 338/26» en la lista de impresos
_RE_ITEM = re.compile(r"(?<![\w/.])(\d{1,2})\s*[.)](?=\s)")


def _numeros_od(lista: str, anio_por_defecto: int) -> list[str]:
    """«340 a 342 y 344» → ['340/26', '341/26', '342/26', '344/26']."""
    piezas = re.findall(rf"({_NUM})|(\bAL?\b|-)", lista)
    salida: list[str] = []
    rango = False
    for numero, guion in piezas:
        if guion:
            rango = bool(salida)
            continue
        n, _, a = numero.replace(" ", "").partition("/")
        anio = anio_cuatro_digitos(a) % 100 if a else anio_por_defecto % 100
        if rango and salida:
            desde = int(salida[-1].split("/")[0])
            if 0 < int(n) - desde <= MAX_RANGO:
                salida += [f"{i}/{anio:02d}" for i in range(desde + 1, int(n) + 1)]
            rango = False
            continue
        salida.append(f"{int(n)}/{anio:02d}")
    return salida


def referencias(texto: str, anio: int, sueltas: bool = False) -> list[tuple[int, Referencia]]:
    """Órdenes del día y expedientes del texto, con su posición, en orden."""
    n = _preparar(texto)
    hallados: list[tuple[int, Referencia]] = []
    for patron in [_RE_OD] + ([_RE_OD_SUELTA] if sueltas else []):
        for m in patron.finditer(n):
            for od in _numeros_od(m.group("lista"), anio):
                hallados.append((m.start(), Referencia("od", od=od)))
    for pos, ident in encontrar(texto):
        hallados.append((pos, Referencia("expediente", ident=ident)))
    vistos, salida = set(), []
    for pos, ref in sorted(hallados, key=lambda x: x[0]):
        if ref.clave not in vistos:
            vistos.add(ref.clave)
            salida.append((pos, ref))
    return salida


def _preparar(texto: str) -> str:
    """Mayúsculas sin tildes, con el mismo largo que el texto (para cortar el original por posición)."""
    return "".join(normalizar(c) or c if c.strip() else " " for c in texto)


def _titulo(texto: str) -> str:
    """El tema sin los números: «OD 409/26 y Anexo - Salud mental» → «Salud mental»."""
    t = re.sub(r"^\s*\d{1,2}\s*[.)]\s+", "", texto)                            # «2. …» de una lista pegada
    for patron in (_RE_OD, _RE_OD_SUELTA, _EN_TEXTO_ORIGEN, _EN_TEXTO_NUMERO):
        n = _preparar(t)
        for m in reversed(list(patron.finditer(n))):
            t = t[:m.start()] + " " + t[m.end():]
    t = re.sub(r"(?i)\(\s*(?:con|y)?\s*anexos?\b[^)]*\)", " ", t)            # «(con anexo)», «( Anexo ; Anexo I )»
    t = re.sub(r"(?i)\by\s+anexos?\b|\banexos?\b", " ", t)                    # «OD 409/26 y Anexo - …»
    t = re.sub(r"\(\s*[;,]?\s*\)", " ", t)
    t = re.sub(r"(?<![\w])[a-z]\)\s*", ", ", t)                                # «a) Cancillería: b) …»
    t = re.sub(r"(?i)\bO\.?D\.?\s+de\s+acuerdos\b", "Acuerdos", t)
    t = limpiar(t)
    t = re.sub(r"\s*,\s*(?=[:,]|$)|(?<=:)\s*,", "", t)
    t = re.sub(r"\s+(?:y|e)\s*$", "", t)
    t = t.strip(" -–:;,.●•")
    if t.startswith("(") and t.endswith(")") and t.count("(") == 1:
        t = t[1:-1]
    if t.count("(") > t.count(")"):
        t += ")"
    if t.startswith(")"):
        t = t[1:]
    return limpiar(re.sub(r"\s+-\s*$|^\s*-\s+", "", t)).strip(" -–:;,")


def _items_numerados(seccion: str) -> list[str]:
    """Corta «1. … 2. … 3) …» respetando la numeración correlativa."""
    marcas, esperado = [], 1
    for m in _RE_ITEM.finditer(seccion):
        if int(m.group(1)) == esperado:
            marcas.append(m)
            esperado += 1
    if len(marcas) < 2 and not (marcas and marcas[0].start() < 5):
        return []
    return [seccion[m.end():(marcas[i + 1].start() if i + 1 < len(marcas) else len(seccion))]
            for i, m in enumerate(marcas)]


def _vinietas(seccion: str) -> list[str]:
    partes = re.split(r"[●•]|\n\s*[-–]\s+|\n", seccion)
    return [p for p in (limpiar(x) for x in partes) if p]


def leer(texto: str, fuente: str = "", anio: int | None = None) -> PlanDeLabor:
    """Asuntos del plan en orden de tratamiento, sin repetir los que figuran en más de una sección."""
    original = unicodedata.normalize("NFKC", texto or "")
    plano = limpiar(original)
    n = _preparar(plano)
    plan = PlanDeLabor([], fuente=fuente, texto=original)
    m = re.search(r"REALIZADA EL DIA (\d{1,2}/\d{1,2}/\d{4})", n)
    plan.fecha_reunion = parsear_fecha(m.group(1)) if m else None
    m = re.search(r"FECHA DE SESION:?\s*(\d{1,2}/\d{1,2}/\d{4})", n)
    plan.fecha_sesion = parsear_fecha(m.group(1)) if m else None
    m = re.search(r"HORA DE SESION:?\s*(\d{1,2}:\d{2})", n)
    plan.hora = m.group(1) if m else ""
    anio = anio or (plan.fecha_sesion or plan.fecha_reunion or date.today()).year

    marcas = []
    for nombre, patron in _SECCIONES:
        for mm in re.finditer(patron, n):
            marcas.append((mm.start(), mm.end(), nombre))
    marcas.sort()
    vistos: set[str] = set()

    def sumar(seccion: str, trozo: str, sueltas: bool = False) -> None:
        trozo = limpiar(trozo)
        if not trozo:
            return
        refs = [r for _p, r in referencias(trozo, anio, sueltas)]
        nuevas = [r for r in refs if r.clave not in vistos]
        if refs and not nuevas:
            return                                  # ya figura en una sección anterior
        if not refs and seccion == "Órdenes del día":
            return                                  # títulos de grupo de la lista de impresos
        vistos.update(r.clave for r in nuevas)
        titulo = _titulo(trozo) if refs else trozo
        if not titulo and re.search(r"\bO\.?D\.? DE ACUERDOS?\b", normalizar(trozo)):
            titulo = "Acuerdos"
        plan.asuntos.append(Asunto(len(plan.asuntos) + 1, titulo, seccion, trozo, nuevas, "ANEXO" in normalizar(trozo)))

    if not any(nombre == "Temas acordados" for _a, _b, nombre in marcas):
        # Lista pegada a mano: una línea por asunto (o todo el texto en orden).
        for linea in [x for x in original.splitlines() if x.strip()] or [plano]:
            sumar("Temas acordados", linea)
        return plan
    for i, (_inicio, fin, nombre) in enumerate(marcas):
        if nombre == "Fin":
            break
        hasta = marcas[i + 1][0] if i + 1 < len(marcas) else len(plano)
        seccion = plano[fin:hasta]
        if nombre == "Temas acordados":
            items = _items_numerados(seccion) or _vinietas(seccion)
            for item in items:
                sumar(nombre, item)
        elif nombre == "Órdenes del día":
            # Lista de impresos: «Grupo ● 337/26 ● 363/26 Otro grupo ● 339/26». Suma las OD que falten y
            # le da el nombre del grupo a los temas acordados que no lo tenían.
            grupo = ""
            for item in re.split(r"(?=[●•])", seccion):
                refs = [r for _p, r in referencias(item, anio, sueltas=True)]
                for ref in refs:
                    for asunto in plan.asuntos:
                        if not asunto.titulo and any(x.clave == ref.clave for x in asunto.referencias) and grupo:
                            asunto.titulo = grupo
                sumar(nombre, item, sueltas=True)
                resto = limpiar(re.sub(r"^[●•]\s*(?:0?D\s*|O\.?D\.?\s*)?\d{1,4}\s*/\s*\d{2,4}", "", limpiar(item)))
                resto = re.sub(r"^\(\s*y?\s*anexos?[^)]*\)", "", resto, flags=re.I).strip(" :")
                if resto and not referencias(resto, anio, sueltas=True):
                    grupo = resto.rstrip(":").strip()
        else:
            for item in _vinietas(seccion):
                sumar(nombre, item)
    return plan


# ======================================================= órdenes del día → expedientes

def expedientes_de_od(html: str) -> list[IdExpediente]:
    """Expedientes que trata una orden del día (página «ordenDelDiaResultadoLink»)."""
    salida: list[IdExpediente] = []
    for a in sopa(html).find_all("a", href=RE_URL_FICHA):
        m = RE_URL_FICHA.search(a["href"])
        ident = IdExpediente(int(m.group("num")), anio_cuatro_digitos(m.group("anio")), m.group("origen").upper(),
                             m.group("tipo").upper())
        if all(ident.corto != x.corto for x in salida):
            salida.append(ident)
    return salida


def url_od(od: str) -> str:
    numero, _, anio = od.partition("/")
    return URL_OD.format(anio=anio_cuatro_digitos(anio), numero=int(numero))


__all__ = ["PlanDeLabor", "Asunto", "Referencia", "PlanPublicado", "ErrorPlan", "leer", "texto_de_archivo",
           "texto_de_pdf", "parsear_plenario", "planes_publicados", "plan_publicado", "referencias",
           "expedientes_de_od", "url_od"]
