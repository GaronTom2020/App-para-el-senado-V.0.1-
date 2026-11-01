"""Escritura del informe de sesión (Word, HTML, Excel) y CSV.

Cada asunto lleva una ficha breve de sus expedientes y, si la sesión ya se hizo, lo que dice la
versión taquigráfica: el desenlace, las votaciones y la «versión revisada» del proyecto con las
modificaciones aprobadas en el recinto en negrita. El resto del articulado remite al dictamen.
"""

from __future__ import annotations

import csv
from html import escape
from pathlib import Path

from . import __version__, plan_labor
from .contenido import nombre_tipo
from .informe_docx import (
    SERIF, _Escritor, _bordes_tabla, _borde_celda, _hipervinculo, _indice_word, _margenes_celdas, _sombrear,
)
from .informe_html import pagina
from .sesion import ResultadoSesion, Tema, comisiones, etiqueta_referencias, resumen
from .util import fmt_fecha

# Columnas del CSV para el tablero de sesiones de Power BI (acordadas el 31/10/2026).
COLUMNAS_CSV = ["Fecha de sesión", "Orden", "Expediente", "OD", "Título", "Tipo", "Comisiones"]
LARGO_EXTRACTO = 260

ESTILOS_SESION = """
.indice td.orden{width:2.5em;color:var(--tinta-suave);font-variant-numeric:tabular-nums}
.indice .sin-numero td{color:var(--tinta-suave)}
.tema{margin-top:56px;padding-top:6px;border-top:3px solid var(--cupula)}
.tema > h2{border:0;margin-top:6px}
.tema .seccion{font-size:13px;color:var(--tinta-suave);margin:10px 0 0}
.desenlace{font-family:var(--serif);font-size:17px;margin:4px 0 12px}
.revisado{border-left:3px solid var(--cupula);padding:2px 0 2px 18px;margin:10px 0 18px}
.revisado h4{margin:16px 0 4px}
.revisado p{margin:0 0 6px}
.revisado strong{font-weight:700;color:var(--tinta)}
td.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.volver{font-size:13px;margin:18px 0 0}
"""


def _e(t) -> str:
    return escape(str(t or ""), quote=True)


def _a(texto: str, url: str | None) -> str:
    return f'<a href="{_e(url)}" target="_blank" rel="noopener">{_e(texto)}</a>' if url else _e(texto)


def _recortar(texto: str, largo: int = LARGO_EXTRACTO) -> str:
    texto = texto or ""
    return texto if len(texto) <= largo else texto[:largo].rsplit(" ", 1)[0] + "…"


def _situaciones(tema: Tema) -> str:
    return "; ".join(f"{r.ficha.id.corto}: {r.situacion.titulo}" for r in tema.resultados) or \
        ("; ".join(tema.errores) if tema.errores else "—")


def _resultado_recinto(tema: Tema) -> str:
    """El desenlace según la versión taquigráfica o, si no la hay, la situación de las fichas."""
    if tema.debate is not None and tema.debate.desenlace:
        return tema.debate.desenlace
    return _situaciones(tema)


def _ods(tema: Tema) -> list[tuple[str, str]]:
    """(«OD 413/26», enlace a la página de la OD) para cada orden del día del asunto."""
    return [(f"OD {od}", plan_labor.url_od(od)) for od in tema.asunto.ordenes_del_dia]


def _pdf_dictamen(tema: Tema, od: str) -> str | None:
    for r in tema.resultados:
        for o in r.ficha.ordenes_del_dia:
            if o.numero == od and (o.url_pdf or o.url):
                return o.url_pdf or o.url
    return None


def _quien(mod) -> str:
    if "Secretario" in mod.quien:
        return "leído por Secretaría"
    return f"propuesto por {mod.quien} y aceptado" if mod.quien else "aceptado en el recinto"


def _tiene_ficha(tema: Tema) -> bool:
    return bool(tema.resultados or tema.debate)


# ===================================================================== HTML

def generar_html(res: ResultadoSesion) -> str:
    plan = res.plan
    partes = ['<div class="membrete"><span>H. Senado de la Nación</span>'
              f"<span>Informe del {res.generado:%d/%m/%Y, %H:%M} h</span></div>",
              '<p class="descripcion" style="margin-top:22px">Informe de sesión</p>',
              f'<h1 class="numero">{_e(fmt_fecha(res.fecha, "Sesión"))}</h1>']
    datos = []
    if plan.hora:
        datos.append(f"Hora de la sesión: {plan.hora} h")
    if plan.fecha_reunion:
        datos.append(f"Plan de labor acordado el {fmt_fecha(plan.fecha_reunion)}")
    if plan.fuente:
        datos.append(_a("Plan de labor (fuente)", plan.fuente) if plan.fuente.startswith("http") else
                     f"Plan de labor: {_e(Path(plan.fuente).name)}")
    if res.taquigrafica is not None:
        datos.append(_a("Versión taquigráfica (PDF)", res.taquigrafica.url))
    if datos:
        partes.append('<p class="enlaces">' + "".join(f"<span>{d}</span>" for d in datos) + "</p>")
    partes.append(f'<section class="situacion"><p class="rotulo">Resumen</p><h2>{_e(resumen(res))}</h2></section>')

    con_ficha = [t for t in res.temas if _tiene_ficha(t)]
    partes.append('<nav class="indice-doc" aria-label="Índice"><p class="rotulo">En este informe</p><ol>'
                  '<li><a href="#asuntos">Asuntos en el orden de tratamiento</a></li>'
                  + "".join(f'<li><a href="#tema-{t.asunto.orden}">{t.asunto.orden}. {_e(_recortar(t.titulo, 90))}'
                            "</a></li>" for t in con_ficha) + "</ol></nav>")

    filas = []
    for t in res.temas:
        a = t.asunto
        clase = "" if a.referencias else ' class="sin-numero"'
        enlace = f'<a href="#tema-{a.orden}">{_e(t.titulo)}</a>' if _tiene_ficha(t) else _e(t.titulo)
        seccion = "" if a.seccion == "Temas acordados" else f"<p class='secundario'>{_e(a.seccion)}</p>"
        ods = ", ".join(_a(texto, url) for texto, url in _ods(t))
        exps = ", ".join(_a(r.ficha.id.corto, r.ficha.url) for r in t.resultados) or \
            _e(", ".join(x.corto for x in t.asunto.expedientes)) or "—"
        filas.append(f"<tr{clase}><td class='orden'>{a.orden}</td><td>{enlace}{seccion}</td>"
                     f"<td>{ods or '—'}</td><td>{exps}</td><td>{_e(_resultado_recinto(t))}</td></tr>")
    columna = "Resultado en el recinto" if res.taquigrafica is not None else "Situación"
    partes.append('<h2 id="asuntos">Asuntos en el orden de tratamiento</h2><div class="tabla"><table class="indice">'
                  "<thead><tr><th>N°</th><th>Asunto</th><th>Orden del día</th><th>Expedientes</th>"
                  f"<th>{columna}</th></tr></thead><tbody>{''.join(filas)}</tbody></table></div>")

    for t in con_ficha:
        partes.append(_tema_html(t))

    partes.append("<footer><p>Informe generado automáticamente con el plan de labor, la versión taquigráfica y los "
                  "datos públicos del sitio del H. Senado de la Nación y de la H. Cámara de Diputados. La versión "
                  "revisada es una lectura automática del texto taquigráfico: antes de usarla en un trámite oficial, "
                  f"verificala en la fuente. Programa Informe de Expedientes, versión {_e(__version__)}.</p>")
    if res.advertencias:
        partes.append("<p>Durante la consulta hubo avisos:</p><ul>"
                      + "".join(f"<li>{_e(x)}</li>" for x in res.advertencias) + "</ul>")
    partes.append("</footer>")
    html = pagina(f"Informe de sesión {fmt_fecha(res.fecha, '')} — H. Senado de la Nación", "".join(partes),
                  indice=False)
    return html.replace("</style>", ESTILOS_SESION + "</style>", 1)


def _tema_html(t: Tema) -> str:
    a = t.asunto
    ods = " · ".join(_a(texto, url) for texto, url in _ods(t))
    partes = [f'<section class="tema" id="tema-{a.orden}"><p class="seccion">Asunto {a.orden} · {_e(a.seccion)}'
              + (f" · {ods}" if ods else "") + f"</p><h2>{_e(t.titulo)}</h2>"]
    if t.resultados:
        partes.append("<h3>Expedientes</h3>")
        filas = "".join(
            f"<tr><td class='docs'>{_a(r.ficha.id.corto, r.ficha.url)}<p class='secundario'>{_e(nombre_tipo(r.ficha))}"
            f"</p></td><td>{_e(_recortar(r.ficha.extracto))}</td><td>{_e(', '.join(comisiones(r)) or '—')}</td>"
            f"<td>{_e(r.situacion.titulo)}</td></tr>" for r in t.resultados)
        partes.append("<div class='tabla'><table><thead><tr><th>Expediente</th><th>Extracto</th><th>Comisiones</th>"
                      f"<th>Situación</th></tr></thead><tbody>{filas}</tbody></table></div>")
    d = t.debate
    if d is not None:
        partes.append("<h3>Tratamiento en el recinto</h3>")
        if d.desenlace:
            partes.append(f'<p class="desenlace">{_e(d.desenlace)}</p>')
        if d.votaciones:
            filas = "".join(f"<tr><td>{_e(v.que.capitalize())}</td><td class='num'>{v.afirmativos}</td>"
                            f"<td class='num'>{v.negativos}</td><td class='num'>{_e(v.abstenciones)}</td>"
                            f"<td class='num'>{_e(v.acta or '—')}</td></tr>" for v in d.votaciones)
            partes.append("<div class='tabla'><table><thead><tr><th>Votación</th><th>Afirmativos</th>"
                          "<th>Negativos</th><th>Abstenciones</th><th>Acta de la sesión</th></tr></thead>"
                          f"<tbody>{filas}</tbody></table></div>")
        partes.append(_revisado_html(t))
    if t.errores:
        partes.append("<p class='secundario'>" + " ".join(_e(x) for x in t.errores) + "</p>")
    partes.append('<p class="volver"><a href="#asuntos">↑ Volver a los asuntos</a></p></section>')
    return "".join(partes)


def _revisado_html(t: Tema) -> str:
    d = t.debate
    dictamenes = [(od, _pdf_dictamen(t, od)) for od in d.ordenes_del_dia]
    remision = ", ".join(_a(f"dictamen de la OD {od}", url) for od, url in dictamenes) or "el dictamen"
    partes = ["<h3>Versión revisada del proyecto</h3>"]
    if not d.modificaciones:
        partes.append(f"<p>La versión taquigráfica no registra modificaciones aprobadas en el recinto: rige el texto "
                      f"del {remision}.</p>")
    else:
        partes.append(f"<p>Texto del {remision} con las modificaciones aprobadas en el recinto, <strong>en negrita"
                      "</strong>. Los artículos que no figuran acá se aprobaron como venían en el dictamen.</p>"
                      '<div class="revisado">')
        for m in d.modificaciones:
            partes.append(f"<h4>{_e(m.articulo or 'Texto incorporado')} <span class='secundario'>({_e(_quien(m))})"
                          "</span></h4>" + "".join(f"<p><strong>{_e(x)}</strong></p>" for x in m.texto.split("\n")))
        partes.append("</div>")
    if d.rechazadas:
        partes.append("<p class='secundario'>Propuestas que no se aceptaron: "
                      + "; ".join(_e(f"{m.articulo or 'sin artículo identificado'} ({m.quien})") for m in d.rechazadas)
                      + ".</p>")
    return "".join(partes)


def escribir_html(res: ResultadoSesion, destino: Path) -> Path:
    Path(destino).write_text(generar_html(res), encoding="utf-8")
    return Path(destino)


# ===================================================================== Word

class _EscritorSesion(_Escritor):
    def __init__(self, res: ResultadoSesion):
        super().__init__(res, encabezado=f"H. Senado de la Nación — Informe de sesión del {fmt_fecha(res.fecha)}")
        self.sesion = res

    def escribir(self, destino: Path, ejemplo: bool = False) -> Path:
        from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
        from docx.shared import Cm

        res, plan = self.sesion, self.sesion.plan
        p = self.parrafo("H. Senado de la Nación", 9, "suave", despues=0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(self.ancho), WD_TAB_ALIGNMENT.RIGHT)
        self._run(p, f"\tInforme del {res.generado:%d/%m/%Y, %H:%M} h", 9, "suave")
        self.parrafo("Informe de sesión", 11, "suave", antes=12, despues=0)
        self.parrafo(fmt_fecha(res.fecha, "Sesión"), 30, "tinta", SERIF, despues=4)
        datos = [x for x in (f"Hora de la sesión: {plan.hora} h" if plan.hora else "",
                             f"Plan de labor acordado el {fmt_fecha(plan.fecha_reunion)}" if plan.fecha_reunion else "")
                 if x]
        if datos or plan.fuente or res.taquigrafica is not None:
            p = self.parrafo(" · ".join(datos), 9.5, "suave", despues=8)
            if plan.fuente.startswith("http"):
                self._run(p, "    ")
                _hipervinculo(p, "Plan de labor (fuente)", plan.fuente, 9.5)
            elif plan.fuente:
                self._run(p, f"    Plan de labor: {Path(plan.fuente).name}", 9.5, "suave")
            if res.taquigrafica is not None:
                self._run(p, "    ")
                _hipervinculo(p, "Versión taquigráfica (PDF)", res.taquigrafica.url, 9.5)

        caja = self.doc.add_table(rows=1, cols=1)
        caja.autofit = False
        _bordes_tabla(caja, color="cupula_tenue")
        _margenes_celdas(caja, 160, 140, 260, 200)
        celda = caja.rows[0].cells[0]
        celda.width = Cm(self.ancho)
        _sombrear(celda, "cupula_tenue")
        _borde_celda(celda, "left", "cupula", 36)
        self._run(celda.paragraphs[0], resumen(res), 12, "tinta", SERIF)

        self.parrafo("Índice", 13, "cupula_oscura", SERIF, antes=16, despues=2)
        _indice_word(self.doc, niveles="1-2")

        self.titulo("Asuntos en el orden de tratamiento")
        anchos = [1.0, 5.4, 2.4, 2.6, 5.2]
        columna = "Resultado en el recinto" if res.taquigrafica is not None else "Situación"
        t = self.tabla(["N°", "Asunto", "Orden del día", "Expedientes", columna], anchos)
        for tema in res.temas:
            fila = self.fila(t, anchos)
            color = None if tema.asunto.referencias else "suave"
            self._run(fila.cells[0].paragraphs[0], str(tema.asunto.orden), 9, "suave")
            self._run(fila.cells[1].paragraphs[0], tema.titulo, 9, color)
            if tema.asunto.seccion != "Temas acordados":
                self._run(fila.cells[1].add_paragraph(), tema.asunto.seccion, 7.5, "suave")
            ods = _ods(tema)
            for j, (texto, url) in enumerate(ods):
                po = fila.cells[2].paragraphs[0] if j == 0 else fila.cells[2].add_paragraph()
                _hipervinculo(po, texto, url, 8.5)
            if not ods:
                self._run(fila.cells[2].paragraphs[0], "—", 8.5, "suave")
            for j, r in enumerate(tema.resultados):
                pe = fila.cells[3].paragraphs[0] if j == 0 else fila.cells[3].add_paragraph()
                _hipervinculo(pe, r.ficha.id.corto, r.ficha.url, 8.5)
            if not tema.resultados:
                self._run(fila.cells[3].paragraphs[0], ", ".join(x.corto for x in tema.asunto.expedientes) or "—",
                          8.5, "suave")
            self._run(fila.cells[4].paragraphs[0], _resultado_recinto(tema), 8.5, color)

        for tema in res.temas:
            if not _tiene_ficha(tema):
                continue
            salto = self.doc.add_paragraph()
            salto.add_run().add_break(WD_BREAK.PAGE)
            self._tema(tema)

        self.parrafo("Informe generado automáticamente con el plan de labor, la versión taquigráfica y los datos "
                     "públicos del sitio del H. Senado de la Nación y de la H. Cámara de Diputados. La versión revisada "
                     "es una lectura automática del texto taquigráfico: antes de usarla en un trámite oficial, "
                     f"verificala en la fuente. Programa Informe de Expedientes, versión {__version__}.",
                     8.5, "suave", antes=22)
        if res.advertencias:
            self.parrafo("Durante la consulta hubo avisos:", 8.5, "suave", antes=4)
            for x in res.advertencias:
                self._run(self.doc.add_paragraph(style="List Bullet"), x, 8.5, "suave")
        self.doc.core_properties.title = f"Informe de sesión del {fmt_fecha(res.fecha)}"
        self.doc.core_properties.author = "Informe de Expedientes"
        self.doc.save(str(destino))
        return Path(destino)

    def _tema(self, tema: Tema) -> None:
        p = self.parrafo(f"Asunto {tema.asunto.orden} · {tema.asunto.seccion}", 9, "suave", despues=0)
        for texto, url in _ods(tema):
            self._run(p, " · ", 9, "suave")
            _hipervinculo(p, texto, url, 9)
        self.doc.add_heading(tema.titulo, level=1)
        if tema.resultados:
            self.doc.add_heading("Expedientes", level=2)
            anchos = [2.4, 7.0, 3.4, 3.8]
            t = self.tabla(["Expediente", "Extracto", "Comisiones", "Situación"], anchos)
            for r in tema.resultados:
                fila = self.fila(t, anchos)
                _hipervinculo(fila.cells[0].paragraphs[0], r.ficha.id.corto, r.ficha.url, 9)
                self._run(fila.cells[0].add_paragraph(), nombre_tipo(r.ficha), 7.5, "suave")
                self._run(fila.cells[1].paragraphs[0], _recortar(r.ficha.extracto), 8.5)
                self._run(fila.cells[2].paragraphs[0], ", ".join(comisiones(r)) or "—", 8.5)
                self._run(fila.cells[3].paragraphs[0], r.situacion.titulo, 8.5)
        d = tema.debate
        if d is not None:
            self.doc.add_heading("Tratamiento en el recinto", level=2)
            if d.desenlace:
                self.parrafo(d.desenlace, 11, "tinta", SERIF, despues=6)
            if d.votaciones:
                anchos = [6.6, 2.4, 2.4, 2.6, 2.6]
                t = self.tabla(["Votación", "Afirmativos", "Negativos", "Abstenciones", "Acta de la sesión"], anchos)
                for v in d.votaciones:
                    fila = self.fila(t, anchos)
                    for i, valor in enumerate((v.que.capitalize(), str(v.afirmativos), str(v.negativos),
                                               v.abstenciones, v.acta or "—")):
                        self._run(fila.cells[i].paragraphs[0], valor, 9)
            self._revisado(tema)
        for error in tema.errores:
            self.parrafo(error, 8.5, "suave", antes=4)

    def _revisado(self, tema: Tema) -> None:
        d = tema.debate
        self.doc.add_heading("Versión revisada del proyecto", level=2)
        dictamenes = [(od, _pdf_dictamen(tema, od)) for od in d.ordenes_del_dia]
        p = self.parrafo("Texto del " if d.modificaciones else
                         "La versión taquigráfica no registra modificaciones aprobadas en el recinto: rige el texto del ",
                         10)
        for k, (od, url) in enumerate(dictamenes):
            if k:
                self._run(p, ", ", 10)
            if url:
                _hipervinculo(p, f"dictamen de la OD {od}", url, 10)
            else:
                self._run(p, f"dictamen de la OD {od}", 10)
        if not dictamenes:
            self._run(p, "dictamen", 10)
        if d.modificaciones:
            self._run(p, " con las modificaciones aprobadas en el recinto, ", 10)
            self._run(p, "en negrita", 10, negrita=True)
            self._run(p, ". Los artículos que no figuran acá se aprobaron como venían en el dictamen.", 10)
            for m in d.modificaciones:
                h = self.doc.add_heading(m.articulo or "Texto incorporado", level=3)
                self._run(h, f"  ({_quien(m)})", 9, "suave")
                for parrafo in m.texto.split("\n"):
                    self.parrafo(parrafo, 10, negrita=True, despues=4)
        else:
            self._run(p, ".", 10)
        if d.rechazadas:
            self.parrafo("Propuestas que no se aceptaron: " + "; ".join(
                f"{m.articulo or 'sin artículo identificado'} ({m.quien})" for m in d.rechazadas) + ".",
                8.5, "suave", antes=4)


def escribir_docx(res: ResultadoSesion, destino: Path) -> Path:
    return _EscritorSesion(res).escribir(Path(destino))


# ===================================================================== Excel

def escribir_xlsx(res: ResultadoSesion, destino: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    normal = Font(name="Arial", size=10)
    titulo = Font(name="Arial", size=14, bold=True, color="114E9A")
    blanco = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    enlace = Font(name="Arial", size=10, color="114E9A", underline="single")
    relleno = PatternFill("solid", fgColor="1A64AB")
    arriba = Alignment(vertical="top", wrap_text=True)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Asuntos"
    hoja["A1"] = f"Informe de sesión del {fmt_fecha(res.fecha)}"
    hoja["A1"].font = titulo
    hoja["A2"] = resumen(res)
    hoja["A2"].font = normal
    columnas = ["Orden", "Asunto", "Sección", "OD", "Expediente", "Tipo", "Extracto", "Situación",
                "Resultado en el recinto", "Comisiones", "Caducidad", "Ficha"]
    anchos = [8, 38, 18, 14, 13, 20, 60, 40, 50, 34, 12, 12]
    for i, (nombre, ancho) in enumerate(zip(columnas, anchos), 1):
        c = hoja.cell(row=4, column=i, value=nombre)
        c.font, c.fill = blanco, relleno
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    col_od = columnas.index("OD") + 1
    fila = 5
    for tema in res.temas:
        filas = tema.resultados or [None]
        desenlace = tema.debate.desenlace if tema.debate is not None else ""
        for r in filas:
            valores = [tema.asunto.orden, tema.titulo, tema.asunto.seccion, ", ".join(tema.asunto.ordenes_del_dia)]
            if r is None:
                valores += [", ".join(x.corto for x in tema.expedientes), "", "", "; ".join(tema.errores), desenlace,
                            "", "", ""]
            else:
                f, s = r.ficha, r.situacion
                valores += [f.id.corto, nombre_tipo(f), f.extracto, s.titulo, desenlace, ", ".join(comisiones(r)),
                            s.caducidad.fecha if s.caducidad else None, "Ver ficha"]
            for i, v in enumerate(valores, 1):
                c = hoja.cell(row=fila, column=i, value=v)
                c.font, c.alignment = normal, arriba
                if columnas[i - 1] == "Caducidad" and v:
                    c.number_format = "DD/MM/YYYY"
            if len(tema.asunto.ordenes_del_dia) == 1:
                c = hoja.cell(row=fila, column=col_od)
                c.hyperlink, c.font = plan_labor.url_od(tema.asunto.ordenes_del_dia[0]), enlace
            if r is not None:
                c = hoja.cell(row=fila, column=len(columnas))
                c.hyperlink, c.font = r.ficha.url, enlace
            fila += 1
    hoja.freeze_panes = "A5"
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(columnas))}{max(fila - 1, 4)}"
    libro.save(str(destino))
    return Path(destino)


# ======================================================================= CSV

def filas_csv(res: ResultadoSesion) -> list[list[str]]:
    """Una fila por asunto, para el tablero de sesiones de Power BI."""
    fecha = fmt_fecha(res.fecha, "")
    filas = []
    for tema in res.temas:
        resultados = tema.resultados
        expedientes = [r.ficha.id.corto for r in resultados] or [x.corto for x in tema.expedientes]
        tipos = list(dict.fromkeys(nombre_tipo(r.ficha) for r in resultados))
        comis = list(dict.fromkeys(c for r in resultados for c in comisiones(r)))
        filas.append([fecha, str(tema.asunto.orden), ", ".join(expedientes), ", ".join(tema.asunto.ordenes_del_dia),
                      tema.titulo, ", ".join(tipos), " / ".join(comis)])
    return filas


def escribir_csv(res: ResultadoSesion, destino: Path) -> Path:
    """Punto y coma, UTF-8 con BOM (Excel y Power BI lo abren con tildes) y fechas dd/mm/aaaa."""
    with open(destino, "w", encoding="utf-8-sig", newline="") as archivo:
        escritor = csv.writer(archivo, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        escritor.writerow(COLUMNAS_CSV)
        escritor.writerows(filas_csv(res))
    return Path(destino)
