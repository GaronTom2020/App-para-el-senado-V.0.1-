"""Informe breve de asuntos entrados (HTML y Word): los asuntos de la revisión, agrupados por tipo."""

from __future__ import annotations

from html import escape
from pathlib import Path

from . import __version__
from .asuntos_entrados import TIPOS, Asunto, ResultadoAsuntos, resumen
from .informe_html import pagina
from .util import fmt_fecha, nombre_archivo_seguro

NOMBRE_TIPO = {"PL": "Proyectos de ley", "PR": "Proyectos de resolución", "PC": "Proyectos de comunicación",
               "PD": "Proyectos de declaración"}
OTROS = "Otros asuntos"


def _e(t) -> str:
    return escape(str(t or ""), quote=True)


def por_tipo(asuntos: list[Asunto]) -> list[tuple[str, list[Asunto]]]:
    """[(«Proyectos de ley», [...]), …] en el orden ley, resolución, comunicación, declaración y otros."""
    grupos: dict[str, list[Asunto]] = {}
    for a in asuntos:
        grupos.setdefault(NOMBRE_TIPO.get(a.tipo, OTROS), []).append(a)
    orden = list(NOMBRE_TIPO.values()) + [OTROS]
    return [(nombre, grupos[nombre]) for nombre in orden if nombre in grupos]


def _titulo(res: ResultadoAsuntos) -> str:
    from datetime import date
    return f"Asuntos entrados desde el {fmt_fecha(date.fromisoformat(res.desde))}" if res.desde else "Asuntos entrados"


def generar_html(res: ResultadoAsuntos, asuntos: list[Asunto], filtros: str = "") -> str:
    partes = ['<div class="membrete"><span>H. Senado de la Nación</span>'
              f"<span>Informe del {res.generado:%d/%m/%Y, %H:%M} h</span></div>",
              f'<h1 class="numero" style="font-size:40px">{_e(_titulo(res))}</h1>',
              f'<p class="extracto">{_e(resumen(res, len(asuntos)))}</p>']
    if filtros:
        partes.append(f'<p class="secundario">Filtros: {_e(filtros)}</p>')
    for nombre, lista in por_tipo(asuntos):
        filas = "".join(
            f"<tr><td class='docs'><a href='{_e(a.url_ficha)}' target='_blank' rel='noopener'>{_e(a.expediente)}</a></td>"
            f"<td class='fecha'>{_e(fmt_fecha(a.fecha_d))}</td><td>{_e(a.extracto)}"
            f"<p class='secundario'>{_e(', '.join(a.giros))}</p></td></tr>" for a in lista)
        partes.append(f"<h2>{_e(nombre)} ({len(lista)})</h2><div class='tabla'><table><thead><tr><th>Expediente</th>"
                      f"<th>Entrada</th><th>Extracto y giro</th></tr></thead><tbody>{filas}</tbody></table></div>")
    if not asuntos:
        partes.append('<p class="vacio">No hay asuntos para mostrar.</p>')
    partes.append("<footer><p>Informe generado con el Diario de Asuntos Entrados (DAE) del sitio del H. Senado de la "
                  f"Nación (DAE leídos: {_e(', '.join(res.daes))}). Sistema de Información Parlamentaria, versión "
                  f"{_e(__version__)}.</p></footer>")
    return pagina(f"{_titulo(res)} — H. Senado de la Nación", "".join(partes))


def escribir_docx(res: ResultadoAsuntos, asuntos: list[Asunto], destino: Path, filtros: str = "") -> Path:
    from docx.enum.section import WD_ORIENT  # noqa: F401 - el escritor arma la página
    from .informe_docx import SERIF, _Escritor, _hipervinculo

    w = _Escritor(res, encabezado=f"H. Senado de la Nación — {_titulo(res)}")
    w.parrafo("H. Senado de la Nación", 9, "suave", despues=0)
    w.parrafo(_titulo(res), 24, "tinta", SERIF, antes=10, despues=4)
    w.parrafo(resumen(res, len(asuntos)), 11, "tinta", SERIF, despues=4)
    if filtros:
        w.parrafo(f"Filtros: {filtros}", 9, "suave", despues=6)
    anchos = [2.6, 2.2, w.ancho - 4.8]
    for nombre, lista in por_tipo(asuntos):
        w.titulo(f"{nombre} ({len(lista)})")
        t = w.tabla(["Expediente", "Entrada", "Extracto y giro"], anchos)
        for a in lista:
            fila = w.fila(t, anchos)
            _hipervinculo(fila.cells[0].paragraphs[0], a.expediente, a.url_ficha, 9)
            w._run(fila.cells[1].paragraphs[0], fmt_fecha(a.fecha_d), 9)
            w._run(fila.cells[2].paragraphs[0], a.extracto, 9)
            if a.giros:
                w._run(fila.cells[2].add_paragraph(), ", ".join(a.giros), 8, "suave")
    if not asuntos:
        w.parrafo("No hay asuntos para mostrar.", 10, "suave", cursiva=True)
    w.parrafo("Informe generado con el Diario de Asuntos Entrados (DAE) del sitio del H. Senado de la Nación "
              f"(DAE leídos: {', '.join(res.daes)}). Sistema de Información Parlamentaria, versión {__version__}.",
              8.5, "suave", antes=18)
    w.doc.core_properties.title = _titulo(res)
    w.doc.save(str(destino))
    return Path(destino)


def escribir(res: ResultadoAsuntos, asuntos: list[Asunto], carpeta: Path, formatos=("docx", "html"),
             filtros: str = "") -> dict[str, Path]:
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    base = nombre_archivo_seguro(f"Asuntos entrados ({res.generado:%Y-%m-%d %H%M})")
    for formato in formatos:
        destino = carpeta / f"{base}.{formato}"
        if formato == "html":
            destino.write_text(generar_html(res, asuntos, filtros), encoding="utf-8")
        elif formato == "docx":
            escribir_docx(res, asuntos, destino, filtros)
        else:
            continue
        res.archivos[formato] = destino
    return res.archivos


__all__ = ["escribir", "generar_html", "escribir_docx", "por_tipo", "TIPOS"]
