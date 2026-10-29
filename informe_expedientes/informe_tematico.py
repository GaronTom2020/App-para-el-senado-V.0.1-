"""Escritura del informe temático en HTML, Word y Excel."""

from __future__ import annotations

import os
import re
from html import escape
from pathlib import Path

from . import __version__, config
from .contenido import Fragmento
from .informe_docx import SERIF, _bordes_tabla, _borde_celda, _Escritor, _hipervinculo, _margenes_celdas, _sombrear
from .informe_html import ESTILOS
from .tematico import ETAPA, Item, ResultadoTematico, resumen, resumen_tipos
from .util import fmt_fecha, limpiar, sin_acentos

COLUMNAS_PRINCIPALES = ["Expediente", "Extracto", "Situación", "Comisión", "Ingreso"]


# ================================================================== comunes

def tramos(extracto: str, palabras: list[str]) -> list[tuple[str, bool]]:
    """Divide el extracto en tramos, marcando las palabras encontradas (sin distinguir tildes)."""
    texto = limpiar(extracto)
    paralelo = sin_acentos(texto).upper()
    if not palabras or len(paralelo) != len(texto):
        return [(texto, False)]
    marcas = []
    for p in palabras:
        patron = re.compile(rf"(?<![A-Z0-9]){re.escape(sin_acentos(limpiar(p)).upper())}[A-Z0-9]*")
        marcas += [(m.start(), m.end()) for m in patron.finditer(paralelo)]
    unidas: list[tuple[int, int]] = []
    for a, b in sorted(marcas):
        if unidas and a <= unidas[-1][1]:
            unidas[-1] = (unidas[-1][0], max(b, unidas[-1][1]))
        else:
            unidas.append((a, b))
    salida, pos = [], 0
    for a, b in unidas:
        if a > pos:
            salida.append((texto[pos:a], False))
        salida.append((texto[a:b], True))
        pos = b
    if pos < len(texto):
        salida.append((texto[pos:], False))
    return salida


def _grupos(res: ResultadoTematico) -> list[tuple[str | None, list[Item]]]:
    grupos: list[tuple[str | None, list[Item]]] = []
    for it in res.items:
        if grupos and grupos[-1][0] == it.nivel:
            grupos[-1][1].append(it)
        else:
            grupos.append((it.nivel, [it]))
    return grupos


def _detalle_situacion(it: Item) -> list[str]:
    lineas = []
    if it.situacion and it.situacion.caducidad:
        c = it.situacion.caducidad
        lineas.append(f"{'Caducidad registrada' if c.registrada else 'Caducidad estimada'}: {fmt_fecha(c.fecha)}")
    if it.reuniones_buscadas:
        if it.reuniones:
            ultima = max((r.fecha for r in it.reuniones if r.fecha), default=None)
            n = len(it.reuniones)
            lineas.append(f"Tratado en {n} {'reunión' if n == 1 else 'reuniones'}"
                          + (f"; la última, el {fmt_fecha(ultima)}" if ultima else ""))
        else:
            lineas.append("Sin reuniones publicadas en que figure")
    if it.error:
        lineas.append(f"No se pudo leer la ficha: {it.error}")
    return lineas


def _metodo(res: ResultadoTematico) -> str:
    texto = ("La búsqueda usa el buscador avanzado del Senado (por extracto) y un filtro propio que no distingue "
             "tildes ni mayúsculas; cada palabra se busca como comienzo de palabra.")
    if res.detallados:
        texto += " La etapa de cada expediente se deduce de su ficha pública."
    if res.busco_reuniones:
        texto += " Las reuniones son las publicadas por las comisiones del giro."
    return texto + " Verificá los datos en la fuente antes de usarlos en un trámite oficial."


# ===================================================================== HTML

ESTILOS_TEMA = """
.antetitulo{font-size:14px;color:var(--tinta-suave);margin:22px 0 0}
.titulo-tema{font-family:var(--serif);font-weight:normal;font-size:40px;line-height:1.12;margin:4px 0 10px}
.resumen{margin:22px 0 8px;padding:20px 26px;background:var(--cupula-tenue);border-left:5px solid var(--cupula)}
.resumen p.frase{font-family:var(--serif);font-size:18px;line-height:1.55;margin:0 0 8px}
.conteos{display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:20px;margin-top:14px}
.conteos table{font-size:13.5px}
.conteos td{padding:5px 8px;border-color:var(--linea)}
.conteos td.n{text-align:right;font-variant-numeric:tabular-nums;width:4em}
.conteos th{background:transparent;padding:4px 8px}
tr.grupo th{background:var(--cupula-tenue);color:var(--cupula-oscura);font-size:14px;padding:9px 10px}
td.exp{white-space:nowrap}
td.exp a{font-weight:600}
td.extracto{font-size:13.5px;line-height:1.45;min-width:22ch}
mark{background:var(--ambar-tenue);color:inherit;padding:0 1px}
"""


def _e(t) -> str:
    return escape(str(t or ""), quote=True)


def _a(texto: str, url: str | None) -> str:
    return f'<a href="{_e(url)}" target="_blank" rel="noopener">{_e(texto)}</a>' if url else _e(texto)


def generar_html(res: ResultadoTematico, carpeta: Path | None = None, ejemplo: bool = False) -> str:
    c = res.consulta
    p: list[str] = []
    a = p.append
    a('<div class="membrete"><span>H. Senado de la Nación</span>'
      f'<span>Informe del {res.generado:%d/%m/%Y, %H:%M} h</span></div>')
    if ejemplo:
        a('<p class="secundario"><strong>Ejemplo generado con páginas de prueba.</strong> '
          'No reemplaza una consulta real.</p>')
    a('<p class="antetitulo">Informe temático</p>')
    a(f'<h1 class="titulo-tema">{_e(c.titulo())}</h1>')

    a('<section class="resumen" aria-label="Resumen">')
    a(f'<p class="frase">{_e(resumen(res))}</p>')
    if res.items:
        a(f'<p class="secundario">{_e(resumen_tipos(res))}</p>')
        a('<div class="conteos">')
        if res.detallados:
            filas = "".join(f'<tr><td>{_e(ETAPA[n])}</td><td class="n">{k}</td></tr>' for n, k in res.conteo_etapas())
            a(f'<table><thead><tr><th>Etapa</th><th></th></tr></thead><tbody>{filas}</tbody></table>')
        filas = "".join(f'<tr><td>{_e(config.ORIGENES.get(o, o))}</td><td class="n">{k}</td></tr>'
                        for o, k in res.conteo(lambda it: it.hallazgo.ident.origen))
        a(f'<table><thead><tr><th>Origen</th><th></th></tr></thead><tbody>{filas}</tbody></table>')
        a("</div>")
    a("</section>")

    a("<h2>Criterios de búsqueda</h2>")
    a('<dl class="datos">' + "".join(f"<dt>{_e(k)}</dt><dd>{_e(v)}</dd>" for k, v in c.criterios()) + "</dl>")

    if res.items:
        a("<h2>Expedientes</h2>")
        columnas = COLUMNAS_PRINCIPALES if res.detallados else COLUMNAS_PRINCIPALES[:2]
        a('<div class="tabla"><table><thead><tr>' + "".join(f"<th>{x}</th>" for x in columnas)
          + "</tr></thead><tbody>")
        for nivel, items in _grupos(res):
            if res.detallados:
                a(f'<tr class="grupo"><th colspan="{len(columnas)}">{_e(ETAPA[nivel])} ({len(items)})</th></tr>')
            for it in items:
                i = it.hallazgo.ident
                tipo = config.TIPOS.get(i.tipo or "", i.tipo or "")
                extracto = "".join(f"<mark>{_e(t)}</mark>" if m else _e(t)
                                   for t, m in tramos(it.hallazgo.extracto, c.palabras))
                fila = (f'<td class="exp">{_a(i.corto, it.url)}<p class="secundario">{_e(tipo)}<br>'
                        f'{_e(config.ORIGENES.get(i.origen, i.origen))}</p></td>'
                        f'<td class="extracto">{extracto}</td>')
                if res.detallados:
                    situacion = _e(it.situacion.titulo if it.situacion else ETAPA[it.nivel])
                    extra = "".join(f'<p class="secundario">{_e(x)}</p>' for x in _detalle_situacion(it))
                    if it.informes.get("html") and carpeta:
                        relativo = os.path.relpath(it.informes["html"], carpeta).replace(os.sep, "/")
                        extra += f'<p class="secundario"><a href="{_e(relativo)}">Informe completo</a></p>'
                    ingreso = fmt_fecha(it.ficha.fecha_mesa_entradas) if it.ficha else "—"
                    fila += (f"<td>{situacion}{extra}</td><td>{_e(it.comisiones() or '—')}</td>"
                             f'<td class="fecha">{ingreso}</td>')
                a(f"<tr>{fila}</tr>")
        a("</tbody></table></div>")

    if res.notas_usuario:
        a(f"<h2>Notas</h2><p class='notas'>{_e(res.notas_usuario)}</p>")
    a("<footer>")
    a(f"<p>{_e(_metodo(res))} Programa Informe de Expedientes, versión {_e(__version__)}.</p>")
    if res.advertencias:
        a("<p>Durante la consulta hubo avisos:</p><ul>" + "".join(f"<li>{_e(x)}</li>" for x in res.advertencias)
          + "</ul>")
    a("</footer>")
    return ('<!DOCTYPE html>\n<html lang="es"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            f"<title>Informe temático: {_e(c.titulo())} — H. Senado de la Nación</title>"
            f"<style>{ESTILOS}{ESTILOS_TEMA}</style></head><body><main class='hoja' style='max-width:1080px'>"
            + "".join(p) + "</main></body></html>\n")


def escribir_html(res: ResultadoTematico, destino: Path, ejemplo: bool = False) -> Path:
    destino = Path(destino)
    destino.write_text(generar_html(res, destino.parent, ejemplo), encoding="utf-8")
    return destino


# ===================================================================== Word

class _EscritorTematico(_Escritor):
    def __init__(self, res: ResultadoTematico):
        super().__init__(res, encabezado=f"H. Senado de la Nación — Informe temático: {res.consulta.titulo()}",
                         apaisado=True)

    def escribir(self, destino: Path, ejemplo: bool = False) -> Path:  # noqa: C901 - documento lineal
        r = self.r
        c = r.consulta
        from docx.enum.text import WD_TAB_ALIGNMENT
        from docx.shared import Cm, Pt

        p = self.parrafo("H. Senado de la Nación", 9, "suave", despues=0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(self.ancho), WD_TAB_ALIGNMENT.RIGHT)
        self._run(p, f"\tInforme del {r.generado:%d/%m/%Y, %H:%M} h", 9, "suave")
        if ejemplo:
            p = self.parrafo("Ejemplo generado con páginas de prueba. ", 9, "lacre", negrita=True, antes=4)
            self._run(p, "No reemplaza una consulta real.", 9, "lacre")
        self.parrafo("Informe temático", 11, "suave", antes=12, despues=0)
        self.parrafo(c.titulo(), 24, "tinta", SERIF, despues=8)

        caja = self.doc.add_table(rows=1, cols=1)
        caja.autofit = False
        _bordes_tabla(caja, color="cupula_tenue")
        _margenes_celdas(caja, 160, 140, 260, 200)
        celda = caja.rows[0].cells[0]
        celda.width = Cm(self.ancho)
        _sombrear(celda, "cupula_tenue")
        _borde_celda(celda, "left", "cupula", 36)
        self._run(celda.paragraphs[0], resumen(r), 12.5, "tinta", SERIF)
        if r.items:
            p = celda.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            self._run(p, resumen_tipos(r), 9.5, "suave")

        if r.items and r.detallados:
            self.titulo("Expedientes por etapa")
            anchos = [7.0, 3.0]
            t = self.tabla(["Etapa", "Expedientes"], anchos)
            for nivel, k in r.conteo_etapas():
                fila = self.fila(t, anchos)
                self._run(fila.cells[0].paragraphs[0], ETAPA[nivel], 9.5)
                self._run(fila.cells[1].paragraphs[0], str(k), 9.5)

        self.titulo("Criterios de búsqueda")
        self.pares([(k, [Fragmento(v)]) for k, v in c.criterios()])

        if r.items:
            self.titulo("Expedientes")
            if r.detallados:
                anchos = [2.9, 10.9, 5.5, 3.9, 2.1]
                columnas = COLUMNAS_PRINCIPALES
            else:
                anchos = [3.0, self.ancho - 3.0]
                columnas = COLUMNAS_PRINCIPALES[:2]
            t = self.tabla(columnas, anchos)
            for nivel, items in _grupos(r):
                if r.detallados:
                    fila = self.fila(t, anchos)
                    unida = fila.cells[0].merge(fila.cells[-1])
                    _sombrear(unida, "cupula_tenue")
                    self._run(unida.paragraphs[0], f"{ETAPA[nivel]} ({len(items)})", 9.5, "cupula_oscura",
                              negrita=True)
                for it in items:
                    i = it.hallazgo.ident
                    fila = self.fila(t, anchos)
                    _hipervinculo(fila.cells[0].paragraphs[0], i.corto, it.url, 9.5, negrita=True)
                    self._run(fila.cells[0].add_paragraph(), f"{i.tipo or '—'}, {i.origen}", 8, "suave")
                    pe = fila.cells[1].paragraphs[0]
                    for texto, marcado in tramos(it.hallazgo.extracto, c.palabras):
                        self._run(pe, texto, 8.5, negrita=marcado)
                    if r.detallados:
                        self._run(fila.cells[2].paragraphs[0],
                                  it.situacion.titulo if it.situacion else ETAPA[it.nivel], 9)
                        for linea in _detalle_situacion(it):
                            self._run(fila.cells[2].add_paragraph(), linea, 8, "suave")
                        self._run(fila.cells[3].paragraphs[0], it.comisiones() or "—", 8.5)
                        self._run(fila.cells[4].paragraphs[0],
                                  fmt_fecha(it.ficha.fecha_mesa_entradas) if it.ficha else "—", 8.5)

        if r.notas_usuario:
            self.titulo("Notas")
            for linea in r.notas_usuario.splitlines():
                self.parrafo(linea, 10.5, fuente=SERIF)
        self.parrafo(_metodo(r) + f" Programa Informe de Expedientes, versión {__version__}.", 8.5, "suave",
                     antes=18)
        if r.advertencias:
            self.parrafo("Durante la consulta hubo avisos:", 8.5, "suave", antes=4)
            for x in r.advertencias:
                self._run(self.doc.add_paragraph(style="List Bullet"), x, 8.5, "suave")
        self.doc.core_properties.title = f"Informe temático: {c.titulo()}"
        self.doc.core_properties.author = "Informe de Expedientes"
        self.doc.save(str(destino))
        return Path(destino)


def escribir_docx(res: ResultadoTematico, destino: Path, ejemplo: bool = False) -> Path:
    return _EscritorTematico(res).escribir(Path(destino), ejemplo)


# ==================================================================== Excel

def escribir_xlsx(res: ResultadoTematico, destino: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    normal = Font(name="Arial", size=10)
    titulo = Font(name="Arial", size=14, bold=True, color="1F4E45")
    negrita = Font(name="Arial", size=10, bold=True)
    blanco = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    enlace = Font(name="Arial", size=10, color="1F4E45", underline="single")
    relleno = PatternFill("solid", fgColor="2E6B5F")
    tenue = PatternFill("solid", fgColor="E4EFEB")
    arriba = Alignment(vertical="top")
    envolver = Alignment(vertical="top", wrap_text=True)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Expedientes"
    columnas = [("Expediente", 12), ("Tipo", 7), ("Origen", 8), ("Extracto", 70), ("Etapa", 20), ("Situación", 38),
                ("Comisión", 28), ("Ingreso", 11), ("Dado cuenta", 11), ("Caducidad", 11), ("Orden del Día", 18),
                ("Reuniones", 10), ("Última reunión", 13), ("Palabras encontradas", 26), ("Ficha", 48)]
    for col, (nombre, ancho) in enumerate(columnas, 1):
        celda = hoja.cell(row=1, column=col, value=nombre)
        celda.font, celda.fill, celda.alignment = blanco, relleno, Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.freeze_panes = "B2"

    for fila, it in enumerate(res.items, 2):
        i, f, s = it.hallazgo.ident, it.ficha, it.situacion
        od = f.ordenes_del_dia[-1] if f and f.ordenes_del_dia else None
        ultima = max((x.fecha for x in it.reuniones if x.fecha), default=None)
        valores = [
            i.corto, i.tipo or "", i.origen, it.hallazgo.extracto, ETAPA[it.nivel],
            s.titulo if s else (it.error or ""), it.comisiones(),
            f.fecha_mesa_entradas if f else None, f.fecha_dado_cuenta if f else None,
            s.caducidad.fecha if s and s.caducidad else None,
            f"{od.numero} ({od.estado.capitalize()})" if od else "",
            len(it.reuniones) if it.reuniones_buscadas else None, ultima,
            ", ".join(it.hallazgo.coincidencias), it.url,
        ]
        for col, valor in enumerate(valores, 1):
            celda = hoja.cell(row=fila, column=col, value=valor)
            celda.font = normal
            celda.alignment = envolver if col in (4, 6, 7) else arriba
            if col in (8, 9, 10, 13) and valor:
                celda.number_format = "DD/MM/YYYY"
        hoja.cell(row=fila, column=1).hyperlink = it.url
        hoja.cell(row=fila, column=1).font = enlace
    ultima_fila = max(2, len(res.items) + 1)
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ultima_fila}"

    r = libro.create_sheet("Resumen")
    r.column_dimensions["A"].width = 30
    r.column_dimensions["B"].width = 44
    r.column_dimensions["C"].width = 14
    r["A1"], r["A1"].font = "Informe temático", titulo
    r["A2"], r["A2"].font = res.consulta.titulo(), negrita
    r["A3"], r["A3"].font = f"Generado el {res.generado:%d/%m/%Y %H:%M} h con datos públicos de senado.gob.ar", normal
    r["A4"], r["A4"].font = resumen(res), normal
    r["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    r.merge_cells("A4:C4")
    r.row_dimensions[4].height = 48
    fila = 6
    for clave, valor in res.consulta.criterios():
        r.cell(row=fila, column=1, value=clave).font = negrita
        r.cell(row=fila, column=2, value=valor).font = normal
        fila += 1

    def tabla(fila: int, titulos: list[str], renglones: list[tuple], columna_hoja: str) -> int:
        for col, t in enumerate(titulos, 1):
            celda = r.cell(row=fila, column=col, value=t)
            celda.font, celda.fill = blanco, relleno
        primera = fila + 1
        rango = f"Expedientes!${columna_hoja}$2:${columna_hoja}${ultima_fila}"
        for renglon in renglones:
            fila += 1
            for col, valor in enumerate(renglon, 1):
                r.cell(row=fila, column=col, value=valor).font = normal
            r.cell(row=fila, column=len(titulos), value=f'=COUNTIF({rango},A{fila})').font = normal
        fila += 1
        r.cell(row=fila, column=1, value="Total").font = negrita
        celda = r.cell(row=fila, column=len(titulos),
                       value=f"=SUM({get_column_letter(len(titulos))}{primera}:{get_column_letter(len(titulos))}{fila - 1})")
        celda.font = negrita
        for col in range(1, len(titulos) + 1):
            r.cell(row=fila, column=col).fill = tenue
        return fila + 2

    fila += 1
    if res.detallados:
        fila = tabla(fila, ["Etapa", "", "Expedientes"], [(ETAPA[n], "") for n, _ in res.conteo_etapas()], "E")
    tipos = sorted({it.hallazgo.ident.tipo or "" for it in res.items})
    fila = tabla(fila, ["Tipo", "Nombre", "Expedientes"], [(t, config.TIPOS.get(t, "")) for t in tipos], "B")
    origenes = sorted({it.hallazgo.ident.origen for it in res.items})
    tabla(fila, ["Origen", "Nombre", "Expedientes"], [(o, config.ORIGENES.get(o, "")) for o in origenes], "C")
    libro.save(str(destino))
    return Path(destino)
